"""Extract embedded drawing images from Excel workbooks.

Excel stores cell images as drawing objects rather than cell values.  The
import pipeline therefore keeps their coordinates and metadata separate from
the tabular values and reads the bytes again only when a reviewed batch is
applied.
"""
from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MIME_BY_FORMAT = {
    "png": "image/png",
    "jpeg": "image/jpeg",
    "jpg": "image/jpeg",
    "gif": "image/gif",
    "webp": "image/webp",
    "bmp": "image/bmp",
    "tiff": "image/tiff",
}
EXTENSION_BY_MIME = {value: f".{key}" for key, value in MIME_BY_FORMAT.items()}
EXTENSION_BY_MIME["image/jpeg"] = ".jpg"


@dataclass(frozen=True)
class ExcelImage:
    sheet_name: str
    source_row: int
    source_column: int
    source_cell: str
    source_field_name: str
    filename: str
    mime_type: str
    width: int | None
    height: int | None
    sha256: str
    content: bytes

    def metadata(self, *, include_hash: bool = True) -> dict[str, Any]:
        result: dict[str, Any] = {
            "sheet_name": self.sheet_name,
            "source_row": self.source_row,
            "source_column": self.source_column,
            "source_cell": self.source_cell,
            "source_field_name": self.source_field_name,
            "filename": self.filename,
            "mime_type": self.mime_type,
            "width": self.width,
            "height": self.height,
        }
        if include_hash:
            result["sha256"] = self.sha256
        return result


def _image_mime(image: Any, content: bytes) -> tuple[str, str]:
    image_format = str(getattr(image, "format", "") or "").lower()
    if image_format == "jpg":
        image_format = "jpeg"
    mime_type = MIME_BY_FORMAT.get(image_format)
    if mime_type:
        return mime_type, EXTENSION_BY_MIME[mime_type]
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png", ".png"
    if content.startswith(b"\xff\xd8\xff"):
        return "image/jpeg", ".jpg"
    if content.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif", ".gif"
    if content.startswith(b"RIFF") and content[8:12] == b"WEBP":
        return "image/webp", ".webp"
    return "application/octet-stream", ".bin"


def _anchor_position(image: Any) -> tuple[int, int] | None:
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return None
    row = getattr(marker, "row", None)
    column = getattr(marker, "col", None)
    if not isinstance(row, int) or not isinstance(column, int):
        return None
    return row, column


def extract_embedded_images(
    file_content: bytes,
    filename: str,
    sheet_name: str | int | None = None,
    columns: list[str] | None = None,
) -> list[ExcelImage]:
    """Return images anchored to data cells in an ``.xlsx`` workbook.

    ``source_row`` follows the import pipeline convention: Excel row 2 is
    source row 2 because row 1 is the header.  Legacy ``.xls`` files and
    delimited text files have no supported drawing extraction and return an
    empty list.
    """
    suffix = Path(filename).suffix.lower()
    if suffix != ".xlsx" or not file_content:
        return []

    workbook = load_workbook(BytesIO(file_content), read_only=False, data_only=False)
    try:
        if sheet_name is None:
            worksheets = list(workbook.worksheets[:1])
        elif isinstance(sheet_name, int):
            worksheets = [workbook.worksheets[sheet_name]]
        else:
            if sheet_name not in workbook.sheetnames:
                return []
            worksheets = [workbook[sheet_name]]

        extracted: list[ExcelImage] = []
        sequence = 0
        for worksheet in worksheets:
            headers = {
                index: str(worksheet.cell(row=1, column=index).value or "").strip()
                for index in range(1, worksheet.max_column + 1)
            }
            for image in getattr(worksheet, "_images", []) or []:
                position = _anchor_position(image)
                if position is None:
                    continue
                anchor_row, anchor_column = position
                # Ignore objects attached to the header row.  Source rows and
                # patent identity resolution start at Excel row 2.
                if anchor_row < 1:
                    continue
                content = image._data()
                mime_type, extension = _image_mime(image, content)
                sequence += 1
                original_name = Path(str(getattr(image, "path", "") or "")).name
                filename_value = original_name or f"image_{sequence}{extension}"
                if "." not in filename_value:
                    filename_value = f"{filename_value}{extension}"
                column_number = anchor_column + 1
                source_row = anchor_row + 1
                source_cell = f"{get_column_letter(column_number)}{source_row}"
                extracted.append(ExcelImage(
                    sheet_name=worksheet.title,
                    source_row=source_row,
                    source_column=anchor_column,
                    source_cell=source_cell,
                    source_field_name=(
                        columns[anchor_column]
                        if columns is not None and anchor_column < len(columns)
                        else headers.get(column_number, "")
                    ),
                    filename=filename_value,
                    mime_type=mime_type,
                    width=getattr(image, "width", None),
                    height=getattr(image, "height", None),
                    sha256=sha256(content).hexdigest(),
                    content=content,
                ))
        return extracted
    finally:
        workbook.close()
